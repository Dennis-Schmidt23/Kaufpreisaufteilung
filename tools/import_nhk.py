import json
from pathlib import Path

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]

EXCEL_DATEI = (
    PROJECT_ROOT
    / "Kaufpreisaufteilung-Grundstuecke-Arbeitshilfe-Berechnung.xlsx"
)

AUSGABE_DATEI = PROJECT_ROOT / "data" / "nhk.json"


def importiere_nhk_daten() -> None:
    workbook = openpyxl.load_workbook(
        EXCEL_DATEI,
        data_only=True,
        read_only=True,
    )

    worksheet = workbook["SW-NHK"]

    daten = []

    for zeile in range(23, 71):
        code = worksheet.cell(zeile, 2).value

        if code is None:
            continue

        datensatz = {
            "code": str(code),
            "gebaeudeart": worksheet.cell(zeile, 3).value or "",
            "beschreibung": worksheet.cell(zeile, 4).value or "",
            "nhk_1": worksheet.cell(zeile, 5).value,
            "nhk_2": worksheet.cell(zeile, 6).value,
            "nhk_3": worksheet.cell(zeile, 7).value,
            "nhk_4": worksheet.cell(zeile, 8).value,
            "nhk_5": worksheet.cell(zeile, 9).value,
            "bgf_wohnflaeche_faktor": worksheet.cell(zeile, 10).value,
            "gesamtnutzungsdauer": worksheet.cell(zeile, 16).value,
        }

        daten.append(datensatz)

    AUSGABE_DATEI.parent.mkdir(exist_ok=True)

    with AUSGABE_DATEI.open("w", encoding="utf-8") as datei:
        json.dump(
            daten,
            datei,
            ensure_ascii=False,
            indent=2,
        )

    print(f"{len(daten)} NHK-Datensätze importiert.")


if __name__ == "__main__":
    importiere_nhk_daten()