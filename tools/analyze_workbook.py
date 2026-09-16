from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "Kaufpreisaufteilung-Grundstuecke-Arbeitshilfe-Berechnung.xlsx"

wb = load_workbook(WORKBOOK, data_only=False)

for sheet in wb.worksheets:
    print("=" * 80)
    print(sheet.title)
    print("=" * 80)

    formulas = 0

    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas += 1
                print(f"{cell.coordinate:6}  {cell.value}")

    print()
    print(f"Anzahl Formeln: {formulas}")
    print()