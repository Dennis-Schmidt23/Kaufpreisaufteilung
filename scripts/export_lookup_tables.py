from pathlib import Path
import json

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

WORKBOOK = ROOT / "Kaufpreisaufteilung-Grundstuecke-Arbeitshilfe-Berechnung.xlsx"
OUTPUT = ROOT / "data"

OUTPUT.mkdir(exist_ok=True)

wb = load_workbook(WORKBOOK, data_only=True)
ws = wb["KPA"]


def read_column(start_row: int, end_row: int, column: int):
    values = []

    for row in range(start_row, end_row + 1):
        value = ws.cell(row=row, column=column).value

        if value is None:
            continue

        values.append(str(value))

    return values


grundstuecksarten = read_column(
    start_row=131,
    end_row=189,
    column=4,      # Spalte D
)

with open(
    OUTPUT / "grundstuecksarten.json",
    "w",
    encoding="utf-8",
) as f:
    json.dump(
        grundstuecksarten,
        f,
        ensure_ascii=False,
        indent=4,
    )

print(f"{len(grundstuecksarten)} Einträge exportiert.")