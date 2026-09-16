from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "Kaufpreisaufteilung-Grundstuecke-Arbeitshilfe-Berechnung.xlsx"

wb = load_workbook(WORKBOOK, data_only=False)

sheet = wb["Fiktives Baujahr"]

print(f"Tabellenblatt: {sheet.title}")
print("-" * 80)

for row in sheet.iter_rows():
    for cell in row:
        if cell.value is None:
            continue

        print(
            f"{cell.coordinate:<6}"
            f"{str(cell.value):<40}"
        )