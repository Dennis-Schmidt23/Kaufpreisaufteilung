from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "Kaufpreisaufteilung-Grundstuecke-Arbeitshilfe-Berechnung.xlsx"

SHEET_NAME = "Fiktives Baujahr"

wb = load_workbook(WORKBOOK, data_only=False)
ws = wb[SHEET_NAME]

print(f"Tabellenblatt: {ws.title}")
print("=" * 100)

for row in ws.iter_rows():
    values = []

    for cell in row:
        if cell.value is None:
            continue

        values.append(f"{cell.coordinate}: {cell.value}")

    if values:
        print(" | ".join(values))